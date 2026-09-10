from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter


DOC_HEADER_TERMS = ("DOCUMENT NO", "DOC. NO", "DOC NO", "DOCUMENT NUMBER", "도서번호", "문서번호")
TITLE_HEADER_TERMS = ("DOCUMENT TITLE", "DOC. TITLE", "도서명", "문서명")
PURPOSE_HEADER_TERMS = ("PURPOSE", "용도")
REV_TERMS = ("REV", "REV.", "REVISION")
SUBMIT_TERMS = ("제출일", "ISSUE DATE", "SUBMISSION DATE", "SUBMIT DATE")
RECEIPT_TERMS = ("접수일", "RETURN DATE", "RECEIPT DATE", "RECEIVED DATE")
STATUS_TERMS = ("STATUS", "APPROVAL CODE", "승인코드")


def norm(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).upper()


def header_match(value: object, terms: tuple[str, ...]) -> bool:
    text = norm(value)
    return any(term in text for term in terms)


def looks_like_doc_no(value: object) -> bool:
    text = str(value or "").strip()
    return (
        8 <= len(text) <= 100
        and text.count("-") >= 2
        and " " not in text
        and bool(re.search(r"[A-Za-z]", text))
        and bool(re.search(r"\d", text))
    )


def merged_value(ws, row: int, col: int) -> tuple[object, int, int]:
    for area in ws.merged_cells.ranges:
        if area.min_row <= row <= area.max_row and area.min_col <= col <= area.max_col:
            return ws.cell(area.min_row, area.min_col).value, area.min_col, area.max_col
    return ws.cell(row, col).value, col, col


def find_target_sheet(wb):
    best = None
    for ws in wb.worksheets:
        for col in range(1, ws.max_column + 1):
            hits = [row for row in range(1, ws.max_row + 1) if looks_like_doc_no(ws.cell(row, col).value)]
            score = len(hits)
            if score and (best is None or score > best[0]):
                best = (score, ws, col, min(hits), max(hits))
    if best is None:
        raise ValueError("No document-number column could be detected.")
    return best


def find_named_column(ws, header_rows: range, terms: tuple[str, ...], fallback: int | None = None):
    for row in header_rows:
        for col in range(1, ws.max_column + 1):
            if header_match(ws.cell(row, col).value, terms):
                return col
    return fallback


def detect_issue_groups(ws, header_rows: range, first_data_row: int):
    candidates = []
    for row in header_rows:
        for col in range(1, ws.max_column + 1):
            text = norm(ws.cell(row, col).value)
            match = re.search(r"\b(\d{1,2})(?:ST|ND|RD|TH)?\s*ISSUE\b", text)
            if match:
                _, start, end = merged_value(ws, row, col)
                candidates.append((int(match.group(1)), row, start, end, str(ws.cell(row, col).value)))
    groups = []
    for issue, issue_row, start, end, label in sorted(set(candidates)):
        group = {"issue": issue, "label": label, "revision": None, "submissionDate": None, "reviewers": {}}
        for col in range(start, end + 1):
            values = [ws.cell(row, col).value for row in header_rows if row > issue_row]
            parent_labels = []
            for upper_row in range(issue_row + 1, first_data_row):
                parent_value, _, _ = merged_value(ws, upper_row, col)
                if parent_value not in (None, ""):
                    parent_labels.append(str(parent_value).strip())
            reviewer_parent = next(
                (
                    value
                    for value in parent_labels
                    if not header_match(value, REV_TERMS + SUBMIT_TERMS + RECEIPT_TERMS + STATUS_TERMS)
                    and norm(value) != "DATE"
                ),
                None,
            )
            if any(header_match(value, REV_TERMS) for value in values):
                group["revision"] = get_column_letter(col)
            if any(header_match(value, SUBMIT_TERMS) for value in values):
                group["submissionDate"] = get_column_letter(col)
            receipt_by_context = (
                any(norm(value) == "DATE" for value in values)
                and reviewer_parent is not None
                and any(term in norm(reviewer_parent) for term in ("RETURN", "OWNER", "REVIEW", "한난"))
            )
            if any(header_match(value, RECEIPT_TERMS) for value in values) or receipt_by_context:
                reviewer = reviewer_parent or f"reviewer_{get_column_letter(col)}"
                group["reviewers"].setdefault(reviewer, {})["receiptDate"] = get_column_letter(col)
            if any(header_match(value, STATUS_TERMS) for value in values):
                reviewer = reviewer_parent or f"reviewer_{get_column_letter(col)}"
                group["reviewers"].setdefault(reviewer, {})["status"] = get_column_letter(col)
        groups.append(group)
    return groups


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook")
    parser.add_argument("--project", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--status", action="append", default=[])
    args = parser.parse_args()

    workbook_path = Path(args.workbook).resolve()
    wb = openpyxl.load_workbook(workbook_path, data_only=False, read_only=False)
    score, ws, doc_col, first_data_row, last_data_row = find_target_sheet(wb)
    header_rows = range(1, first_data_row)
    detected_doc_col = find_named_column(ws, header_rows, DOC_HEADER_TERMS, doc_col)
    title_col = find_named_column(ws, header_rows, TITLE_HEADER_TERMS)
    purpose_col = find_named_column(ws, header_rows, PURPOSE_HEADER_TERMS)
    issues = detect_issue_groups(ws, header_rows, first_data_row)

    required_headers = []
    for row in header_rows:
        for col in range(1, ws.max_column + 1):
            value = ws.cell(row, col).value
            if value not in (None, ""):
                required_headers.append({"cell": f"{get_column_letter(col)}{row}", "value": str(value)})
    merged_ranges = sorted(str(area) for area in ws.merged_cells.ranges)
    signature_source = json.dumps(
        {"sheet": ws.title, "headers": required_headers, "merges": merged_ranges},
        ensure_ascii=False,
        sort_keys=True,
    ).encode("utf-8")

    profile = {
        "project": args.project,
        "workbookExample": str(workbook_path),
        "sheetName": ws.title,
        "sheetNames": wb.sheetnames,
        "usedRange": ws.calculate_dimension(),
        "headerSignature": hashlib.sha256(signature_source).hexdigest(),
        "headerRows": [1, first_data_row - 1],
        "requiredHeaders": required_headers,
        "mergedRanges": merged_ranges,
        "documentNumberColumn": get_column_letter(detected_doc_col),
        "documentTitleColumn": get_column_letter(title_col) if title_col else None,
        "purposeColumn": get_column_letter(purpose_col) if purpose_col else None,
        "firstDataRow": first_data_row,
        "lastDataRow": last_data_row,
        "documentCount": score,
        "allowedStatuses": args.status,
        "issues": issues,
    }
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(profile, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "project": profile["project"],
        "sheetName": profile["sheetName"],
        "documentNumberColumn": profile["documentNumberColumn"],
        "dataRows": [first_data_row, last_data_row],
        "documentCount": score,
        "issueCount": len(issues),
        "headerSignature": profile["headerSignature"],
        "output": str(output.resolve()),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
